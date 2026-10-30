"""
メンバー一覧ファイル（CSV / Excel）の読み取り。

ここでは「ファイルを行×列の文字列表に直す」ところまでを担当する。
どの列が名前・学年・スキルなのかの解釈と検証は、ブラウザ側の取り込みプレビュー
（列の紐づけ画面）でユーザーが確認しながら行う。
"""

import csv
import io

from openpyxl import load_workbook

# 取り込み対象として受け付ける最大サイズ・最大行数・最大列数
MAX_IMPORT_FILE_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_IMPORT_COLUMNS = 30

CSV_EXTENSIONS = (".csv", ".tsv", ".txt")
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")


def decode_csv_bytes(raw):
    """
    CSV のバイト列を文字列にする。(text, encoding名) を返す。

    Excel が「CSV UTF-8」で保存すると BOM 付き UTF-8、「CSV（コンマ区切り）」で保存すると
    Shift_JIS(cp932) になるため、この2つを順に試す。
    """
    for encoding in ("utf-8-sig", "cp932"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8（一部の文字を置換）"


def _detect_delimiter(text):
    """先頭の数行から区切り文字（, / タブ / ;）を推定する。"""
    sample_lines = [line for line in text.splitlines() if line.strip()][:5]
    if not sample_lines:
        return ","
    counts = {delimiter: sum(line.count(delimiter) for line in sample_lines) for delimiter in (",", "\t", ";")}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def _cell_to_text(value):
    """セルの値を表示用の文字列にする（3.0 は "3"、空は ""）。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _trim_rows(rows):
    """各行の末尾の空セルを落とし、列数の上限を適用する。"""
    trimmed = []
    for row in rows:
        cells = [_cell_to_text(cell) for cell in row][:MAX_IMPORT_COLUMNS]
        while cells and cells[-1] == "":
            cells.pop()
        trimmed.append(cells)
    return trimmed


def parse_csv(raw):
    text, encoding = decode_csv_bytes(raw)
    delimiter = _detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for row in reader:
        rows.append(row)
        if len(rows) > MAX_IMPORT_ROWS:
            break
    truncated = len(rows) > MAX_IMPORT_ROWS
    rows = _trim_rows(rows[:MAX_IMPORT_ROWS])
    return {
        "encoding": encoding,
        "sheets": [{"name": "CSV", "rows": rows, "truncated": truncated}],
    }


def parse_excel(raw):
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl は壊れたファイルで多様な例外を出す
        raise ValueError(f"Excelファイルを開けませんでした（{exc}）") from exc

    sheets = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            truncated = False
            for row in worksheet.iter_rows(values_only=True):
                if len(rows) >= MAX_IMPORT_ROWS:
                    truncated = True
                    break
                rows.append(row)
            sheets.append({"name": worksheet.title, "rows": _trim_rows(rows), "truncated": truncated})
    finally:
        workbook.close()

    if not sheets:
        raise ValueError("Excelファイルにシートがありません")
    return {"encoding": None, "sheets": sheets}


def parse_members_file(filename, raw):
    """
    アップロードされたファイルを読み、{"encoding", "sheets": [{"name", "rows", "truncated"}]} を返す。
    読めない・対応外の場合は ValueError（メッセージは画面にそのまま表示できる文面）。
    """
    if not raw:
        raise ValueError("ファイルが空です")
    if len(raw) > MAX_IMPORT_FILE_BYTES:
        raise ValueError(f"ファイルが大きすぎます（{MAX_IMPORT_FILE_BYTES // (1024 * 1024)}MBまで）")

    lowered = (filename or "").lower()
    if lowered.endswith(EXCEL_EXTENSIONS):
        result = parse_excel(raw)
    elif lowered.endswith(CSV_EXTENSIONS):
        result = parse_csv(raw)
    elif lowered.endswith(".xls"):
        raise ValueError("古い形式（.xls）には対応していません。Excelで「.xlsx」または「CSV」として保存し直してください")
    else:
        raise ValueError("対応しているファイルは .csv / .xlsx です")

    if all(not any(cell for cell in row) for sheet in result["sheets"] for row in sheet["rows"]):
        raise ValueError("ファイルに読み取れるデータがありません")
    return result
