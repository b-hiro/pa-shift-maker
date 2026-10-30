"""
メンバー一覧ファイル（CSV / Excel）の読み取りと、属性(tag)項目の検証・出力を確認する。
列の紐づけ・検証はブラウザ側で行うため、ここではサーバー側の「ファイル→行×列の表」変換を見る。
"""

import io

import pytest
from openpyxl import Workbook

from app import app, normalize_member, validate_member_structure
from member_import import MAX_IMPORT_FILE_BYTES, parse_members_file


CSV_TEXT = "名前,学年,属性,卓,ステージ\nみらい,1,,3,4\nげんき,おたすけ,お助け枠,1,3\n"


def test_csv_utf8_with_bom():
    raw = CSV_TEXT.encode("utf-8-sig")
    result = parse_members_file("members.csv", raw)
    assert result["encoding"] == "utf-8-sig"
    rows = result["sheets"][0]["rows"]
    assert rows[0] == ["名前", "学年", "属性", "卓", "ステージ"]
    # 空セルは "" として列位置が保たれ、末尾の空セルだけが落ちる
    assert rows[1] == ["みらい", "1", "", "3", "4"]
    assert rows[2][2] == "お助け枠"


def test_csv_shift_jis():
    raw = CSV_TEXT.encode("cp932")
    result = parse_members_file("members.csv", raw)
    assert result["encoding"] == "cp932"
    assert result["sheets"][0]["rows"][2][0] == "げんき"


def test_csv_tab_delimiter():
    raw = "名前\t学年\nみらい\t1\n".encode("utf-8")
    rows = parse_members_file("members.tsv", raw)["sheets"][0]["rows"]
    assert rows[1] == ["みらい", "1"]


def _xlsx_bytes(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_multiple_sheets_and_number_cells():
    raw = _xlsx_bytes({
        "メンバー": [["名前", "学年", "属性", "卓", "ステージ"], ["みらい", 1, None, 3, 4.0]],
        "メモ": [["関係ないシート"]],
    })
    result = parse_members_file("members.xlsx", raw)
    assert [s["name"] for s in result["sheets"]] == ["メンバー", "メモ"]
    # 数値セルは "1" / "4" のように整数表記の文字列になる（"1.0" にはならない）
    assert result["sheets"][0]["rows"][1] == ["みらい", "1", "", "3", "4"]


@pytest.mark.parametrize("filename, raw, message", [
    ("members.csv", b"", "空"),
    ("members.xls", b"dummy", ".xls"),
    ("members.pdf", b"dummy", ".csv / .xlsx"),
    ("members.csv", b"\n\n,,\n", "データがありません"),
    ("members.xlsx", b"not an excel file", "開けませんでした"),
])
def test_unreadable_files_raise_value_error(filename, raw, message):
    with pytest.raises(ValueError) as exc_info:
        parse_members_file(filename, raw)
    assert message in str(exc_info.value)


def test_too_large_file_is_rejected():
    with pytest.raises(ValueError):
        parse_members_file("members.csv", b"a" * (MAX_IMPORT_FILE_BYTES + 1))


def test_parse_endpoint_returns_rows_and_errors():
    client = app.test_client()

    response = client.post(
        "/api/parse-members-file",
        data={"file": (io.BytesIO(CSV_TEXT.encode("utf-8-sig")), "members.csv")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    body = response.get_json()
    assert body["filename"] == "members.csv"
    assert body["sheets"][0]["rows"][1][0] == "みらい"

    # ファイル未指定 / 非対応形式は 400 とメッセージ
    assert client.post("/api/parse-members-file", data={}).status_code == 400
    bad = client.post(
        "/api/parse-members-file",
        data={"file": (io.BytesIO(b"x"), "members.pdf")},
        content_type="multipart/form-data",
    )
    assert bad.status_code == 400
    assert "error" in bad.get_json()


def test_sample_template_csv_is_readable():
    from pathlib import Path

    path = Path(__file__).resolve().parent.parent / "members_template.csv"
    result = parse_members_file(path.name, path.read_bytes())
    rows = result["sheets"][0]["rows"]
    assert rows[0][0] == "名前"
    assert any(row and row[2] == "お助け枠" for row in rows[1:])


def test_excel_summary_has_tag_column():
    from app import create_excel_workbook

    members = [
        {"name": "げんき", "grade": "おたすけ", "tag": "お助け枠", "count": 2, "desk_count": 1, "stage_count": 1},
        {"name": "みらい", "grade": None, "tag": None, "count": 3, "desk_count": 2, "stage_count": 1},
    ]
    wb = create_excel_workbook({}, {}, members)
    ws = wb["全体集計"]
    assert [c.value for c in ws[1]][:3] == ["メンバー名", "学年", "属性"]
    # シフト回数の降順に並ぶ: みらい(3) → げんき(2)
    assert [ws.cell(row=2, column=3).value, ws.cell(row=3, column=3).value] == ["-", "お助け枠"]
    assert ws.cell(row=2, column=4).value == 3


def test_tag_is_optional_string_and_defaults_to_none():
    base = {"name": "A", "skill_desk": 3, "skill_stage": 3, "count": 0}
    assert validate_member_structure(base, 0) is None
    assert validate_member_structure({**base, "tag": "お助け枠"}, 0) is None
    assert validate_member_structure({**base, "tag": None}, 0) is None
    assert validate_member_structure({**base, "tag": 5}, 0) is not None
    assert normalize_member(base)["tag"] is None
    assert normalize_member({**base, "tag": "お助け枠"})["tag"] == "お助け枠"
