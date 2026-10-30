"""
Excel 出力の「個人別一覧」シート（自分の出番だけにフィルターで絞り込むためのシート）のテスト。
"""

import io

from openpyxl import load_workbook

from app import (
    app,
    create_excel_workbook,
    generate_pa_shift_multi_day,
    generate_timetable_multi_day,
)


def _member(name, skill_desk=3, skill_stage=3):
    return {
        "name": name,
        "skill_desk": skill_desk,
        "skill_stage": skill_stage,
        "count": 0,
        "grade": None,
        "ng_bands": [],
        "ng_times": {},
        "req_bands": [],
    }


def _build():
    timetable_multi = generate_timetable_multi_day({
        "num_days": 2,
        "days": [
            {"day_number": 1, "start_time": "11:00", "bands": ["A", "B", "C"], "rh_mins": 10, "act_mins": 10},
            {"day_number": 2, "start_time": "11:00", "bands": ["D", "E"], "rh_mins": 10, "act_mins": 10,
             "break_duration": 30, "break_after_band": "D"},
        ],
    })
    members = [_member(f"L{i}") for i in range(5)] + [_member("A1", 1, 1), _member("Nobody", 1, 1)]
    config = {"desk_max_members": 2, "stage_max_members": 2, "candidate_count": 4}
    shift, updated, infeasible_days = generate_pa_shift_multi_day(timetable_multi, members, config)
    assert not infeasible_days
    return timetable_multi, shift, updated


def test_personal_sheet_is_second_and_has_filter():
    timetable_multi, shift, members = _build()
    wb = create_excel_workbook(timetable_multi, shift, members)

    assert wb.sheetnames[:4] == ["全体集計", "個人別一覧", "1日目", "2日目"]
    ws = wb["個人別一覧"]
    assert [c.value for c in ws[1]] == ["メンバー名", "日", "時間帯", "種別", "バンド名"]  # 役割列は無い
    assert ws.auto_filter.ref == f"A1:E{ws.max_row}"
    assert ws.freeze_panes == "A2"


def test_personal_sheet_has_one_row_per_person_per_entry():
    """1人×(リハ・本番それぞれ)×担当バンドで1行。休憩は含まず、名前は連結されない。"""
    timetable_multi, shift, members = _build()
    ws = create_excel_workbook(timetable_multi, shift, members)["個人別一覧"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]

    # 期待行数: シフト結果の延べ担当数 × 2（リハと本番）
    expected = sum(
        len(team["卓"]) + len(team["ステージ"])
        for day in shift.values() for team in day.values()
    ) * 2
    assert len(rows) == expected

    names = {r[0] for r in rows}
    assert all("、" not in name for name in names)
    assert all(r[3] in ("リハ", "本番") for r in rows)


def test_personal_sheet_matches_summary_counts_and_is_time_ordered():
    timetable_multi, shift, members = _build()
    wb = create_excel_workbook(timetable_multi, shift, members)
    ws = wb["個人別一覧"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]

    for member in members:
        mine = [r for r in rows if r[0] == member["name"]]
        # 本番の行数 = シフト回数
        assert sum(1 for r in mine if r[3] == "本番") == member["count"]
        # 日→時間の昇順（同じ日の中で時間帯が前から順に並ぶ）
        keys = [(int(r[1].replace("日目", "")), r[2]) for r in mine]
        assert keys == sorted(keys)


def test_personal_sheet_follows_registration_order():
    timetable_multi, shift, members = _build()
    ws = create_excel_workbook(timetable_multi, shift, members)["個人別一覧"]
    order = []
    for row in ws.iter_rows(min_row=2):
        name = row[0].value
        if name not in order:
            order.append(name)
    registered = [m["name"] for m in members if m["name"] in order]
    assert order == registered


def test_download_endpoint_includes_personal_sheet():
    timetable_multi, shift, members = _build()
    client = app.test_client()
    res = client.post("/api/download-excel", json={
        "timetable_multi": timetable_multi, "shift": shift, "members": members,
    })
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.data))
    assert "個人別一覧" in wb.sheetnames
